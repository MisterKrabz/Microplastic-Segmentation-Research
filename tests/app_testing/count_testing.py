"""
Count Testing Script
====================
Runs the same YOLO detection pipeline as `tests/model_testing/run_sequential.py`
over the testing images and writes the per-image count into the
"Counted # by v3" column of `datasets/testing_datasets/Quantification/Counting.xlsx`
(matched on the "Image name" column).

The count value is `len(yolo_boxes)` after cluster suppression -- identical to
what `run_sequential.py` reports as `pred_count`. SAM2 is intentionally not run
here because it does not affect the count.
"""

import os
import sys
from typing import Optional

import cv2
import numpy as np
import torch
import openpyxl
from ultralytics import YOLO


# ==========================================
# CONFIGURATION
# ==========================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

SOURCE_PATHS = os.path.normpath(
    os.path.join(SCRIPT_DIR, "./../../datasets/testing_datasets/OriginalImage")
)
YOLO_MODEL_PATH = os.path.normpath(
    os.path.join(SCRIPT_DIR, "./../../models/hunter-yolo-v0.5.4/hunter-yolo-v0.5.4.pt")
)
COUNTING_XLSX = os.path.normpath(
    os.path.join(SCRIPT_DIR, "./../../datasets/testing_datasets/Quantification/Counting.xlsx")
)

CONFIDENCE = 0.2
IOU_THRESH = 0.5
IMG_SIZE = 1280

CLUSTER_CONTAIN_THRESH = 0.85
CLUSTER_MIN_CHILDREN = 2
ENABLE_CLUSTER_SUPPRESSION = True

VALID_EXTS = (".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff")
EXPORT_ROOT_NAME = "_exports"

DEVICE = (
    "cuda" if torch.cuda.is_available()
    else "mps" if torch.backends.mps.is_available()
    else "cpu"
)

IMAGE_NAME_HEADER = "Image name"
V3_COLUMN_HEADER = "Counted # by v3"


# ==========================================
# Image discovery
# ==========================================
def get_image_list(source_paths_str: str) -> list[str]:
    image_paths: list[str] = []
    for dataset_root in (p.strip() for p in source_paths_str.split(",") if p.strip()):
        if os.path.isfile(dataset_root):
            if dataset_root.lower().endswith(VALID_EXTS):
                image_paths.append(dataset_root)
            continue

        if os.path.isdir(dataset_root):
            for r, dirs, files in os.walk(dataset_root):
                dirs[:] = [d for d in dirs if d != EXPORT_ROOT_NAME]
                for f in files:
                    if f.lower().endswith(VALID_EXTS):
                        image_paths.append(os.path.join(r, f))

    return sorted(image_paths)


# ==========================================
# Cluster suppression (same as run_sequential)
# ==========================================
def suppress_cluster_boxes(
    xyxy: np.ndarray,
    contain_thresh: float = CLUSTER_CONTAIN_THRESH,
    min_children: int = CLUSTER_MIN_CHILDREN,
) -> np.ndarray:
    """Return indices of boxes to keep, dropping any box that nearly contains
    `min_children` or more strictly-smaller boxes.
    """
    n = len(xyxy)
    if n <= 1:
        return np.arange(n, dtype=int)

    x1 = xyxy[:, 0]
    y1 = xyxy[:, 1]
    x2 = xyxy[:, 2]
    y2 = xyxy[:, 3]
    areas = np.maximum(0.0, x2 - x1) * np.maximum(0.0, y2 - y1)

    ix1 = np.maximum(x1[:, None], x1[None, :])
    iy1 = np.maximum(y1[:, None], y1[None, :])
    ix2 = np.minimum(x2[:, None], x2[None, :])
    iy2 = np.minimum(y2[:, None], y2[None, :])
    iw = np.clip(ix2 - ix1, 0.0, None)
    ih = np.clip(iy2 - iy1, 0.0, None)
    inter = iw * ih

    area_j = areas[None, :]
    with np.errstate(divide="ignore", invalid="ignore"):
        ios_j = np.where(area_j > 0, inter / area_j, 0.0)

    larger = areas[:, None] > areas[None, :]
    contained = (ios_j >= contain_thresh) & larger
    np.fill_diagonal(contained, False)

    children_count = contained.sum(axis=1)
    keep = children_count < min_children
    return np.nonzero(keep)[0]


# ==========================================
# YOLO inference (same flags as run_sequential)
# ==========================================
def predict_count(yolo: YOLO, img_path: str) -> int:
    results = yolo.predict(
        source=img_path,
        conf=CONFIDENCE,
        iou=IOU_THRESH,
        imgsz=IMG_SIZE,
        verbose=False,
        agnostic_nms=True,
        augment=True,
        max_det=1000,
    )
    r0 = results[0]
    if r0.boxes is None or len(r0.boxes) == 0:
        return 0

    all_xyxy = r0.boxes.xyxy.cpu().numpy().astype(np.float32)
    if ENABLE_CLUSTER_SUPPRESSION:
        keep_idx = suppress_cluster_boxes(all_xyxy)
        return int(len(keep_idx))
    return int(len(all_xyxy))


# ==========================================
# Spreadsheet helpers
# ==========================================
def find_header_row(ws, image_name_header: str) -> int:
    """Search the first ~20 rows for the cell whose value matches the image-name
    header. Returns the 1-indexed row number, or -1 if not found.
    """
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=False):
        for c in row:
            if c.value is not None and str(c.value).strip() == image_name_header:
                return c.row
    return -1


def find_column_for_header(ws, header_row: int, header_text: str) -> int:
    for c in ws[header_row]:
        if c.value is not None and str(c.value).strip() == header_text:
            return c.column
    return -1


def build_name_to_row(ws, header_row: int, name_col: int) -> dict[str, int]:
    name_to_row: dict[str, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=name_col).value
        if v is None:
            continue
        name_to_row[str(v).strip()] = r
    return name_to_row


# ==========================================
# Main
# ==========================================
def main() -> int:
    print(f"Device: {DEVICE}")
    print(f"Source images: {SOURCE_PATHS}")
    print(f"YOLO weights:  {YOLO_MODEL_PATH}")
    print(f"Counting xlsx: {COUNTING_XLSX}")

    if not os.path.exists(COUNTING_XLSX):
        print(f"❌ Missing counting xlsx: {COUNTING_XLSX}")
        return 1
    if not os.path.exists(YOLO_MODEL_PATH):
        print(f"❌ Missing YOLO weights: {YOLO_MODEL_PATH}")
        return 1

    image_files = get_image_list(SOURCE_PATHS)
    if not image_files:
        print(f"❌ No images found at {SOURCE_PATHS}")
        return 1
    print(f"Found {len(image_files)} image(s) to process.\n")

    wb = openpyxl.load_workbook(COUNTING_XLSX)
    ws = wb.active

    header_row = find_header_row(ws, IMAGE_NAME_HEADER)
    if header_row < 0:
        print(f"❌ Could not find '{IMAGE_NAME_HEADER}' header in the first 20 rows.")
        return 1

    name_col = find_column_for_header(ws, header_row, IMAGE_NAME_HEADER)
    v3_col = find_column_for_header(ws, header_row, V3_COLUMN_HEADER)
    if name_col < 0:
        print(f"❌ '{IMAGE_NAME_HEADER}' column not found in row {header_row}.")
        return 1
    if v3_col < 0:
        print(f"❌ '{V3_COLUMN_HEADER}' column not found in row {header_row}.")
        return 1

    print(
        f"Header row: {header_row} | "
        f"name col: {openpyxl.utils.get_column_letter(name_col)} | "
        f"v3 col: {openpyxl.utils.get_column_letter(v3_col)}\n"
    )

    name_to_row = build_name_to_row(ws, header_row, name_col)

    print(f"Loading YOLO: {YOLO_MODEL_PATH}")
    yolo = YOLO(YOLO_MODEL_PATH)
    print()

    total = len(image_files)
    written = 0
    no_match: list[str] = []
    errors: list[tuple[str, str]] = []

    for i, img_path in enumerate(image_files, start=1):
        stem = os.path.splitext(os.path.basename(img_path))[0]

        bgr = cv2.imread(img_path)
        if bgr is None:
            print(f"[{i}/{total}] {stem}: ⚠️  cv2 could not read image, skipping")
            errors.append((stem, "cv2.imread returned None"))
            continue

        try:
            count = predict_count(yolo, img_path)
        except Exception as e:
            print(f"[{i}/{total}] {stem}: ❌ inference error: {e}")
            errors.append((stem, repr(e)))
            continue

        target_row: Optional[int] = name_to_row.get(stem)
        if target_row is None:
            print(f"[{i}/{total}] {stem}: count={count} (no matching row in xlsx)")
            no_match.append(stem)
            continue

        ws.cell(row=target_row, column=v3_col, value=count)
        written += 1
        print(f"[{i}/{total}] {stem}: count={count} -> row {target_row}")

    wb.save(COUNTING_XLSX)

    print()
    print("=" * 60)
    print(f"Wrote {written} count(s) to '{V3_COLUMN_HEADER}' "
          f"(column {openpyxl.utils.get_column_letter(v3_col)}) in:")
    print(f"  {COUNTING_XLSX}")
    if no_match:
        print(f"\n{len(no_match)} image(s) had no matching row in the xlsx:")
        for n in no_match:
            print(f"  - {n}")
    if errors:
        print(f"\n{len(errors)} image(s) failed:")
        for n, msg in errors:
            print(f"  - {n}: {msg}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
